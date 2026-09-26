import json
from datetime import datetime, timezone
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor
from fastapi.testclient import TestClient
from app import db, direct_links
from app.main import app, ist
from app.auth import password_hash
from app.search import save_job
from app.resume import _replace_in_docx, _pdf_with_original_layout
from conftest import csrf
from test_workflow import sample_job


def test_candidate_cannot_reach_owner_activity_or_reset_account(client):
    db.write("UPDATE profile SET name='Surendran Nansen R' WHERE id=1")
    response=client.post('/candidate-account',data={'_csrf':csrf(client),'password':'Ab1234'})
    assert response.status_code==200
    with TestClient(app) as c:
        assert c.post('/login',data={'username':'Surendran Nansen R','password':'Ab1234'}).status_code==200
        home=c.get('/').text
        assert 'Welcome, Surendran Nansen R' in home
        assert '/activity' not in home
        assert c.get('/activity').status_code==403
        assert c.post('/candidate-account',data={'_csrf':csrf(c),'password':'123456'}).status_code==403
        assert 'Surendran’s sign-in' not in c.get('/profile').text
        client.post('/candidate-account',data={'_csrf':csrf(client),'password':'XY1234'})
        assert c.get('/activity',follow_redirects=False).status_code==303


def test_short_password_rate_limit(test_env):
    db.write("UPDATE profile SET name='Person' WHERE id=1")
    db.write('INSERT INTO candidate_account VALUES(1,?,?)',(password_hash('aB1234'),db.utcnow()))
    with TestClient(app) as c:
        for _ in range(8):
            assert c.post('/login',data={'username':'Person','password':'bad'}).status_code==401
        assert c.post('/login',data={'username':'Person','password':'aB1234'}).status_code==429


def test_memory_edits_survive_resume_replacement_and_questions_roll(client,monkeypatch):
    from app.config import settings
    monkeypatch.setattr('app.questions.settings',type('Settings',(),{'openrouter_api_key':''})())
    object.__setattr__(settings,'openrouter_api_key','')
    job=save_job(sample_job('Research Director','Lab','memory'))
    ids=[]
    for i in range(5):
        ids.append(db.write('INSERT INTO questions(job_id,question,requirement,created_at) VALUES(?,?,?,?)',
            (job['id'],f'Useful question {i}?',f'Experience topic {i}',db.utcnow())))
    page=client.get('/jobs/'+job['id']).text
    assert 'Useful question 0?' in page and 'Useful question 3?' not in page
    for i in range(3):
        page=client.post(f'/jobs/{job["id"]}/questions/{ids[i]}',data={'_csrf':csrf(client),'answer':'Yes','detail':f'Example {i}'}).text
    assert 'Thanks for sharing your information.' in page and 'Useful question 3?' in page
    fact=db.one('SELECT * FROM facts ORDER BY id LIMIT 1')
    client.post(f'/facts/{fact["id"]}/edit',data={'_csrf':csrf(client),'fact':'Led a team of ten engineers at my previous employer.'})
    doc=Document();doc.add_paragraph('This is a replacement resume with more than eighty characters. It describes research work, leadership and electrical engineering.')
    stream=BytesIO();doc.save(stream)
    client.post('/resume',data={'_csrf':csrf(client)},files={'resume':('new.docx',stream.getvalue())})
    assert db.one('SELECT fact FROM facts WHERE id=?',(fact['id'],))['fact'].startswith('Led a team')
    client.post(f'/facts/{fact["id"]}/remove',data={'_csrf':csrf(client)})
    assert db.one('SELECT answer FROM questions WHERE id=?',(ids[0],))['answer']=='Removed'


def test_search_filters_only_on_home_and_ist(client):
    assert 'name="role_labels"' not in client.get('/profile').text
    assert 'name="roles"' in client.get('/').text
    assert ist('2026-09-25T21:50:00+00:00')=='26 Sep 2026, 03:20 AM IST'


def test_logo_tracking_and_opened_state(client,monkeypatch):
    monkeypatch.setattr('app.questions.ensure_questions',lambda job:None)
    raw=sample_job('R&D Head','Company','tracked');raw['employer_logo']='https://example.com/logo.png'
    job=save_job(raw)
    run=db.write('INSERT INTO search_runs(kind,started_at,parameters_json) VALUES(?,?,?)',('manual',db.utcnow(),'{}'))
    db.write('INSERT INTO search_results(run_id,job_id,rank,internal_score) VALUES(?,?,1,90)',(run,job['id']))
    assert 'Unopened' in client.get(f'/runs/{run}').text
    client.get('/jobs/'+job['id'])
    assert 'Opened' in client.get(f'/runs/{run}').text
    page=client.get('/my-jobs').text
    assert 'Jobs Opened Status' in page and 'logo.png' in page and 'Times opened' not in page
    assert 'UTC' not in page
    assert db.one('SELECT status FROM job_status WHERE job_id=?',(job['id'],)) is None


def employer_html(title='Research Director',closed=False):
    posting={'@type':'JobPosting','title':title,'hiringOrganization':{'name':'Example Engineering'},
             'jobLocation':{'address':{'addressLocality':'Chennai','addressCountry':'IN'}},
             'validThrough':'2099-01-01'}
    return '<html><script type="application/ld+json">'+json.dumps(posting)+'</script><h1>'+title+'</h1><p>'+('This position is closed' if closed else 'Apply')+'</p></html>'


def test_direct_gate_rejects_boards_wrong_job_and_closed_vacancies(test_env,monkeypatch):
    job={'title':'Research Director','company':'Example Engineering','location':'Chennai','company_url':'https://example.com'}
    assert direct_links.verify_page('https://careers.example.com/jobs/123',employer_html(),job)[0]
    assert not direct_links.verify_page('https://linkedin.com/jobs/123',employer_html(),job)[0]
    assert not direct_links.verify_page('https://careers.example.com/jobs/123',employer_html('Sales Manager'),job)[0]
    assert not direct_links.verify_page('https://careers.example.com/jobs/123',employer_html(closed=True),job)[0]
    assert not direct_links.public_url('https://127.0.0.1/private')
    assert not direct_links.public_url('http://169.254.169.254/')
    assert not direct_links.public_url('https://user:pass@example.com/')


def test_aggregator_redirect_never_recorded_as_application_click(client,monkeypatch):
    job=save_job(sample_job('Research Director','Example Engineering','blocked'))
    monkeypatch.setattr(direct_links,'read_page',lambda url:('https://linkedin.com/jobs/1',employer_html()))
    response=client.get('/out/'+job['id'],follow_redirects=False)
    assert response.status_code==303 and response.headers['location'].startswith('/jobs/')
    assert db.one("SELECT COUNT(*) n FROM job_activity WHERE action='application_link_opened'")['n']==0


def test_web_discovery_still_requires_matching_employer_page(test_env,monkeypatch):
    job=save_job(sample_job('Research Director','Example Engineering','web-discovery'))
    job['company_url']='https://example.com'
    job['raw_json']='{}'
    monkeypatch.setattr(direct_links,'read_page',lambda url:(url,employer_html('Sales Manager')))
    result=direct_links.resolve_job(job,force=True,web_discover=lambda _:['https://example.com/jobs/123'])
    assert result['direct_status']=='unverified' and not result['direct_url']
    monkeypatch.setattr(direct_links,'read_page',lambda url:(url,employer_html()))
    result=direct_links.resolve_job(job,force=True,web_discover=lambda _:['https://example.com/jobs/123'])
    assert result['direct_status']=='verified'
    assert result['direct_url']=='https://example.com/jobs/123'


def test_docx_change_preserves_unchanged_run_styles():
    document=Document();p=document.add_paragraph();r=p.add_run('R&D leadership: ');r.bold=True;r.font.size=Pt(14)
    r=p.add_run('Led electrical research teams.');r.font.name='Cambria';r.font.color.rgb=RGBColor.from_string('123456')
    assert _replace_in_docx(document,[{'original':'Led electrical research teams.','replacement':'Led electrical development teams.'}])
    assert p.runs[0].bold and p.runs[0].font.size==Pt(14)
    assert p.runs[1].font.name=='Cambria' and str(p.runs[1].font.color.rgb)=='123456'
    assert p.text=='R&D leadership: Led electrical development teams.'


def test_pdf_unchanged_export_identical(test_env):
    import pymupdf as fitz
    source=test_env/'source.pdf';doc=fitz.open();page=doc.new_page();page.insert_text((60,60),'Original design');doc.save(source)
    target=test_env/'copy.pdf';_pdf_with_original_layout(source,target,[])
    assert source.read_bytes()==target.read_bytes()


def test_pdf_word_text_remains_editable_and_reupload_readable(test_env):
    import pymupdf as fitz
    from app.pdf_word_layout import pdf_to_word
    from app.resume import extract_text
    source=test_env/'source.pdf';doc=fitz.open();page=doc.new_page()
    wording='Led electrical research and development teams across multiple product lines with responsibility for design and quality.'
    page.insert_text((40,70),wording,fontsize=9);doc.save(source)
    output=test_env/'layout.docx';pdf_to_word(source,output)
    text,_=extract_text(output.read_bytes(),output.name)
    assert wording in text
    document=Document(output)
    assert _replace_in_docx(document,[{'original':'electrical research','replacement':'electrical design'}])
    assert any('electrical design' in (n.text or '') for n in document.element.xpath('.//w:t'))


def test_search_displays_progress_and_preserves_pause(client,monkeypatch):
    from app.config import settings
    object.__setattr__(settings,'jsearch_api_key','test-key')
    calls=[]
    monkeypatch.setattr('app.main.run_search',lambda **kw:calls.append(kw))
    response=client.post('/search',data={'_csrf':csrf(client),'roles':'R&D Head','keywords':'design',
        'locations':'Chennai','work_types':'Consulting','count':'5','days_recent':'7'})
    assert response.status_code==200 and 'Finding and checking company vacancies' in response.text
    assert len(calls)==1
    client.post('/search',data={'_csrf':csrf(client),'count':'5','days_recent':'7'})
    assert len(calls)==1  # Double submit does not start another paid search.
