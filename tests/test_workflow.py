import json
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from docx import Document

from app import db, search
from app.alerts import run_due_alerts
from app.config import settings
from app.resume import apply_changes, export_variant
from conftest import csrf


def sample_job(title, company, suffix):
    return {
        "job_id": suffix, "job_title": title, "employer_name": company,
        "job_location": "Chennai, Tamil Nadu, India", "job_country": "IN",
        "job_employment_type": "Full-time", "job_description":
        "Lead research and development, quality systems and product improvement.",
        "job_apply_link": f"https://example.com/apply/{suffix}",
        "job_posted_at_datetime_utc": "2026-09-24T08:00:00Z",
    }


def test_search_deduplicates_and_tracks_click_separately(client, monkeypatch):
    object.__setattr__(settings, "jsearch_api_key", "test-key")
    object.__setattr__(settings, "openrouter_api_key", "")
    monkeypatch.setattr(search, "fetch_page", lambda *args: (
        [sample_job("R&D Director", "Textile Lab", "a"),
         sample_job("R&D Director", "Textile Lab", "a"),
         sample_job("Junior Assistant", "Other", "b")], None))
    run_id = search.run_search(kind="manual", roles="R&D Director", keywords="quality",
                               locations="Chennai", work_types="Full-time", count=10)
    rows = db.all_rows("SELECT j.title,r.rank FROM search_results r JOIN jobs j ON j.id=r.job_id "
                       "WHERE r.run_id=? ORDER BY r.rank", (run_id,))
    assert len(rows) == 2
    assert rows[0]["title"] == "R&D Director"
    job_id = db.one("SELECT id FROM jobs WHERE title='R&D Director'")["id"]
    response = client.get(f"/out/{job_id}", follow_redirects=False)
    assert response.status_code == 303
    assert response.headers["location"] == "https://example.com/apply/a"
    assert "R&amp;D Director" in client.get("/my-jobs").text
    assert db.one("SELECT status FROM job_status WHERE job_id=?", (job_id,)) is None
    client.post(f"/jobs/{job_id}/status", data={"_csrf": csrf(client), "status": "Applied"})
    assert db.one("SELECT status FROM job_status WHERE job_id=?", (job_id,))["status"] == "Applied"


def test_answer_becomes_editable_profile_fact(client):
    raw = sample_job("Research Advisor", "Life Lab", "c")
    job = search.save_job(raw)
    qid = db.write("INSERT INTO questions(job_id,question,requirement,created_at) VALUES(?,?,?,?)",
                   (job["id"], "Have you led pilot production?", "Led pilot production",
                    db.utcnow()))
    response = client.post(f"/jobs/{job['id']}/questions/{qid}",
                           data={"_csrf": csrf(client), "answer": "Yes",
                                 "detail": "At HCL for five years"}, follow_redirects=False)
    assert response.status_code == 303
    assert "Led pilot production" in client.get("/profile").text
    client.post(f"/jobs/{job['id']}/questions/{qid}",
                data={"_csrf": csrf(client), "answer": "No", "detail": ""})
    assert db.one("SELECT COUNT(*) AS n FROM facts WHERE active=1")["n"] == 0


def test_resume_export_requires_exact_existing_wording(test_env):
    source = test_env / "resume.docx"
    document = Document()
    document.add_paragraph("Led textile research programs.")
    document.save(source)
    original = "Led textile research programs."
    changes = [{"original": original,
                "replacement": "Led textile research programs and pilot trials."}]
    docx_path, pdf_path = export_variant(source, original, changes,
                                         "A Person", "Example Company", 1)
    assert docx_path.name == "A Person - Example Company.docx"
    assert pdf_path.exists() and pdf_path.stat().st_size > 100
    assert "pilot trials" in Document(docx_path).paragraphs[0].text
    try:
        apply_changes(original, [{"original": "Experience never written", "replacement": "X"}])
        assert False, "An unsupported edit should fail"
    except ValueError:
        pass


def test_alert_due_at_ist_once_per_day(test_env, monkeypatch):
    aid = db.write(
        "INSERT INTO alerts(name,email,role_labels,keywords,locations,work_types,time_ist,count,"
        "days_recent,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
        ("Daily", "someone@example.com", "R&D", "", "Chennai", "Full-time", "08:00", 5, 7,
         db.utcnow()),
    )
    sent = []
    from app import alerts
    def fake_run(alert):
        sent.append(alert["id"])
        db.write("UPDATE alerts SET last_sent_date_ist=? WHERE id=?", ("2026-09-25", aid))
        return 3
    monkeypatch.setattr(alerts, "run_alert", fake_run)
    now = datetime(2026, 9, 25, 8, 0, tzinfo=ZoneInfo("Asia/Kolkata"))
    assert run_due_alerts(now) == [(aid, "sent 3")]
    assert run_due_alerts(now) == []
    assert sent == [aid]
