"""Editable Word text positioned on the PDF's original page canvas.

PDFs have no Word paragraph semantics. Fixed page-relative text preserves their
headings, columns, fonts and line breaks without guessing a new document layout.
"""
from io import BytesIO
import pymupdf as fitz
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

V='urn:schemas-microsoft-com:vml'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def pdf_to_word(source,output):
    pdf=fitz.open(source); doc=Document()
    normal=doc.styles['Normal'];normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.space_before=Pt(0)
    for i,page in enumerate(pdf):
        if i:
            from docx.enum.section import WD_SECTION_START
            section=doc.add_section(WD_SECTION_START.NEW_PAGE)
        else:section=doc.sections[0]
        section.page_width=Pt(page.rect.width);section.page_height=Pt(page.rect.height)
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Pt(0)
        section.header_distance=section.footer_distance=Pt(0)
        holder=doc.add_paragraph();holder.paragraph_format.line_spacing=Pt(1)
        # Original graphics form a background; all text is native editable Word text.
        backdrop=fitz.open();backdrop.insert_pdf(pdf,from_page=i,to_page=i);bp=backdrop[0]
        for block in bp.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                for span in line['spans']:
                    bp.add_redact_annot(span['bbox'],fill=False,cross_out=False)
        bp.apply_redactions(images=0,graphics=0)
        bitmap=bp.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).tobytes('png')
        run=holder.add_run();inline=run.add_picture(BytesIO(bitmap),width=Pt(page.rect.width),height=Pt(page.rect.height))._inline
        anchor=OxmlElement('wp:anchor')
        for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'1','layoutInCell':'1','allowOverlap':'1'}.items():anchor.set(k,v)
        simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
        for direction in ('H','V'):
            pos=OxmlElement('wp:position'+direction);pos.set('relativeFrom','page')
            off=OxmlElement('wp:posOffset');off.text='0';pos.append(off);anchor.append(pos)
        for tag in ('extent','effectExtent'):
            el=inline.find(qn('wp:'+tag))
            if el is not None:anchor.append(el)
        anchor.append(OxmlElement('wp:wrapNone'))
        for tag in ('docPr','cNvGraphicFramePr'):
            el=inline.find(qn('wp:'+tag))
            if el is not None:anchor.append(el)
        anchor.append(inline.find(qn('a:graphic')));inline.getparent().replace(inline,anchor)
        index=0
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                spans=[s for s in line['spans'] if s['text'].strip()]
                if not spans:continue
                # Individual spans preserve blue bullets and bold phrases, including columns.
                for span in spans:
                    index+=1;x0,y0,x1,y1=span['bbox']
                    pict=OxmlElement('w:pict');shape=etree.SubElement(pict,'{'+V+'}rect')
                    shape.set('id',f'page{i+1}text{index}');shape.set('type','#_x0000_t202')
                    shape.set('style',f'position:absolute;margin-left:{x0:.3f}pt;margin-top:{y0:.3f}pt;width:{x1-x0+20:.3f}pt;height:{y1-y0+20:.3f}pt;z-index:10;mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
                    shape.set('stroked','f');shape.set('filled','f')
                    box=etree.SubElement(shape,'{'+V+'}textbox');box.set('inset','0pt,0pt,0pt,0pt');box.set('style','mso-fit-shape-to-text:t')
                    content=etree.SubElement(box,'{'+W+'}txbxContent');p=OxmlElement('w:p');content.append(p)
                    pp=OxmlElement('w:pPr');p.append(pp)
                    spacing=OxmlElement('w:spacing');spacing.set(qn('w:before'),'0');spacing.set(qn('w:after'),'0');spacing.set(qn('w:line'),str(round((y1-y0)*20)));spacing.set(qn('w:lineRule'),'exact');pp.append(spacing)
                    r=OxmlElement('w:r');p.append(r);rp=OxmlElement('w:rPr');r.append(rp)
                    fonts=OxmlElement('w:rFonts');family=span['font'].split('+')[-1].replace('-BoldItalic','').replace('-Bold','').replace('-Italic','')
                    for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),family)
                    rp.append(fonts)
                    for tag in ('sz','szCs'):
                        size=OxmlElement('w:'+tag);size.set(qn('w:val'),str(round(span['size']*2)));rp.append(size)
                    color=OxmlElement('w:color');color.set(qn('w:val'),f"{span['color']:06X}");rp.append(color)
                    if span['flags']&16:rp.append(OxmlElement('w:b'))
                    if span['flags']&2:rp.append(OxmlElement('w:i'))
                    text=OxmlElement('w:t');text.set('{http://www.w3.org/XML/1998/namespace}space','preserve');text.text=span['text'];r.append(text)
                    holder.add_run()._r.append(pict)
        backdrop.close()
    doc.save(output);pdf.close()
